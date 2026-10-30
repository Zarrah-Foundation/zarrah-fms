"""Official resolution document rendering (shared context + PDF/DOCX bytes)."""

from __future__ import annotations

import io
import re
from pathlib import Path
from typing import Any

from django.db.models import Prefetch
from django.utils import timezone
from django.utils.html import strip_tags
from django.utils.translation import get_language, gettext as _

from core.models import SiteSettings
from governance.models import Resolution, ResolutionSignature, ResolutionSigner, ResolutionStatus
from governance.services.audit import log_governance

_EMAIL_RE = re.compile(
    r"[a-zA-Z0-9._%+-]+@[a-zA-Z0-9.-]+\.[a-zA-Z]{2,}",
    re.IGNORECASE,
)
_PHONE_RE = re.compile(
    r"(?<!\w)(?:\+?\d{1,3}[\s\-.]?)?(?:\(?\d{2,4}\)?[\s\-.]?)?\d{3,4}[\s\-.]?\d{3,4}(?:[\s\-.]?\d+)?(?!\w)"
)


def _scrub_sensitive_text(text: str) -> str:
    """Remove emails and phone-like sequences from plain text."""
    cleaned = _EMAIL_RE.sub("", text or "")
    cleaned = _PHONE_RE.sub("", cleaned)
    cleaned = re.sub(r"[ \t]{2,}", " ", cleaned)
    cleaned = re.sub(r"\n{3,}", "\n\n", cleaned)
    return cleaned.strip()


def html_to_plain_paragraphs(html: str) -> list[str]:
    """Convert rich-text HTML to sanitized plain-text paragraphs."""
    if not html or not str(html).strip():
        return []
    raw = str(html)
    try:
        import bleach

        plain = bleach.clean(raw, tags=[], strip=True)
    except ImportError:
        plain = strip_tags(raw)
    plain = plain.replace("\r\n", "\n").replace("\r", "\n")
    plain = _scrub_sensitive_text(plain)
    if not plain:
        return []
    parts = re.split(r"\n\s*\n+", plain)
    return [part.strip() for part in parts if part.strip()]


def _pick_bilingual(*, lang: str, bn: str, en: str) -> str:
    bn = (bn or "").strip()
    en = (en or "").strip()
    if lang == "bn":
        return bn or en
    return en or bn


def _format_datetime(dt) -> str:
    if not dt:
        return ""
    return timezone.localtime(dt).strftime("%d %B %Y, %H:%M")


def _safe_meeting_venue(meeting) -> str:
    """Offline venue only — never expose online credentials or passwords."""
    bits: list[str] = []
    if meeting.venue_name:
        bits.append(meeting.venue_name.strip())
    if meeting.room_details:
        bits.append(meeting.room_details.strip())
    if meeting.venue_address:
        bits.append(meeting.venue_address.strip())
    return ", ".join(bits)


def _signature_rows(version) -> list[dict[str, Any]]:
    # Prefer prefetched signers (avoids N+1 on document generation).
    cache = getattr(version, "_prefetched_objects_cache", None) or {}
    if "signers" in cache:
        signers = sorted(
            (s for s in version.signers.all() if s.is_required),
            key=lambda s: (s.selected_at or timezone.now(), s.pk),
        )
    else:
        signers = list(
            ResolutionSigner.objects.filter(version=version, is_required=True)
            .select_related("signature", "member")
            .order_by("selected_at", "pk")
        )
    rows: list[dict[str, Any]] = []
    for signer in signers:
        signature = getattr(signer, "signature", None)
        image_path = None
        if signature and signature.signature_image:
            path = Path(signature.signature_image.path)
            if path.is_file():
                image_path = path
        if signature:
            rows.append(
                {
                    "name": signature.signer_name,
                    "designation": signature.signer_designation,
                    "member_id": signature.member_id_snapshot,
                    "signed_at": signature.signed_at,
                    "signed_at_display": _format_datetime(signature.signed_at),
                    "image_path": image_path,
                    "is_signed": True,
                }
            )
        else:
            member = signer.member
            name = f"{member.first_name} {member.last_name}".strip() or member.member_id
            rows.append(
                {
                    "name": name,
                    "designation": "",
                    "member_id": member.member_id,
                    "signed_at": None,
                    "signed_at_display": "",
                    "image_path": None,
                    "is_signed": False,
                }
            )
    return rows


def build_resolution_document_context(resolution) -> dict[str, Any]:
    """
    Build a sanitized, presentation-ready context dict for PDF/DOCX renderers.

    Excludes meeting passwords, online join details, internal version notes,
    and scrubs emails/phones from body text. Bangla is primary; English is
    optional and omitted when empty.
    """
    resolution = (
        Resolution.objects.select_related(
            "meeting__meeting_type",
            "agenda_item",
            "current_version",
            "finalized_by",
        )
        .prefetch_related(
            Prefetch(
                "current_version__signatures",
                queryset=ResolutionSignature.objects.select_related("signer").order_by(
                    "signed_at", "pk"
                ),
            ),
            Prefetch(
                "current_version__signers",
                queryset=ResolutionSigner.objects.select_related(
                    "member", "signature"
                ).order_by("selected_at", "pk"),
            ),
        )
        .get(pk=resolution.pk)
    )
    site = SiteSettings.load()
    lang = (get_language() or "bn").split("-")[0].lower()
    meeting = resolution.meeting
    version = resolution.current_version
    if version is None:
        raise ValueError("Resolution has no current version.")

    foundation = _pick_bilingual(
        lang=lang,
        bn=site.foundation_name_bn,
        en=site.foundation_name_en or site.site_title,
    )
    foundation_alt = _pick_bilingual(
        lang="en" if lang == "bn" else "bn",
        bn=site.foundation_name_bn or site.short_name,
        en=site.foundation_name_en,
    )

    logo_path = None
    if site.site_logo:
        candidate = Path(site.site_logo.path)
        if candidate.is_file():
            logo_path = candidate

    title_bn = (version.title_bn or resolution.title_bn or "").strip()
    title_en = (version.title_en or resolution.title_en or "").strip()
    resolution_title = title_bn or title_en

    content_paragraphs_bn = html_to_plain_paragraphs(version.content_bn or "")
    content_paragraphs_en = html_to_plain_paragraphs(version.content_en or "")
    decision_paragraphs_bn = html_to_plain_paragraphs(version.decision_bn or "")
    decision_paragraphs_en = html_to_plain_paragraphs(version.decision_en or "")

    # Backward-compatible combined paragraphs (BN first, then EN if distinct).
    content_paragraphs = list(content_paragraphs_bn)
    for para in content_paragraphs_en:
        if para not in content_paragraphs:
            content_paragraphs.append(para)
    decision_paragraphs = list(decision_paragraphs_bn)
    for para in decision_paragraphs_en:
        if para not in decision_paragraphs:
            decision_paragraphs.append(para)

    agenda_title = ""
    agenda_number = None
    if resolution.agenda_item_id:
        agenda = resolution.agenda_item
        agenda_title = (agenda.title_bn or agenda.title_en or "").strip()
        agenda_number = agenda.number

    meeting_title = meeting.localized_title(lang)
    meeting_type = meeting.meeting_type.localized_name(lang)
    meeting_starts = timezone.localtime(meeting.starts_at)
    venue = _safe_meeting_venue(meeting)

    status_display = resolution.get_status_display()
    version_label = _("Version %(num)s") % {"num": version.version_number}
    if resolution.status == ResolutionStatus.FINALIZED and resolution.finalized_at:
        status_line = _("Finalized on %(date)s") % {
            "date": _format_datetime(resolution.finalized_at)
        }
    elif version.is_frozen:
        status_line = _("Version %(num)s — frozen") % {"num": version.version_number}
    else:
        status_line = _("Version %(num)s — %(status)s") % {
            "num": version.version_number,
            "status": status_display,
        }

    generated_at = timezone.localtime(timezone.now())
    document_title = _("Resolution") + f" — {resolution.resolution_id}"
    signatures = _signature_rows(version)
    signed_count = sum(1 for s in signatures if s["is_signed"])
    required_count = len(signatures)

    return {
        "resolution": resolution,
        "version": version,
        "meeting": meeting,
        "site_settings": site,
        "language": lang,
        "foundation_name": foundation,
        "foundation_name_alt": foundation_alt,
        "registration_number": (site.registration_number or "").strip(),
        "website_url": (site.website_url or "").strip(),
        "logo_path": logo_path,
        "document_title": document_title,
        "resolution_id": resolution.resolution_id,
        "resolution_title": resolution_title,
        "title_bn": title_bn,
        "title_en": title_en,
        "resolution_status": resolution.status,
        "resolution_status_display": status_display,
        "version_number": version.version_number,
        "version_label": version_label,
        "version_status_line": status_line,
        "is_finalized": resolution.status == ResolutionStatus.FINALIZED,
        "finalized_at": resolution.finalized_at,
        "finalized_at_display": _format_datetime(resolution.finalized_at),
        "meeting_reference": meeting.reference_id,
        "meeting_title": meeting_title,
        "meeting_type": meeting_type,
        "meeting_starts_at": meeting_starts,
        "meeting_starts_at_display": meeting_starts.strftime("%d %B %Y, %H:%M"),
        "meeting_mode": meeting.get_mode_display(),
        "meeting_venue": venue,
        "agenda_title": agenda_title,
        "agenda_number": agenda_number,
        "content_paragraphs": content_paragraphs,
        "content_paragraphs_bn": content_paragraphs_bn,
        "content_paragraphs_en": content_paragraphs_en,
        "decision_paragraphs": decision_paragraphs,
        "decision_paragraphs_bn": decision_paragraphs_bn,
        "decision_paragraphs_en": decision_paragraphs_en,
        "signatures": signatures,
        "signed_count": signed_count,
        "required_signer_count": required_count,
        "generated_at": generated_at,
        "generated_at_display": generated_at.strftime("%d %B %Y, %H:%M"),
    }


def _log_document_generation(resolution, *, actor, fmt: str, version_number: int) -> None:
    log_governance(
        actor,
        f"RESOLUTION_DOCUMENT_{fmt.upper()}",
        "Resolution",
        resolution.pk,
        f"Generated {fmt.upper()} for {resolution.resolution_id} v{version_number}",
        new={"format": fmt, "version": version_number},
    )


def generate_resolution_pdf(resolution, *, actor=None) -> bytes:
    """Render an official resolution PDF with Bengali-capable fonts and page numbers."""
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import (
        Image,
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    )

    from core.pdf_fonts import ensure_pdf_fonts, escape_pdf_text, pdf_paragraph_styles

    ensure_pdf_fonts()
    ctx = build_resolution_document_context(resolution)
    styles = pdf_paragraph_styles()
    buffer = io.BytesIO()

    def _draw_page_number(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 8)
        canvas.setFillColor(colors.HexColor("#666666"))
        page_text = _("Page %(page)d") % {"page": doc.page}
        canvas.drawCentredString(A4[0] / 2, 10 * mm, page_text)
        footer = ctx["generated_at_display"]
        canvas.drawRightString(A4[0] - 18 * mm, 10 * mm, footer)
        canvas.restoreState()

    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        leftMargin=18 * mm,
        rightMargin=18 * mm,
        topMargin=16 * mm,
        bottomMargin=18 * mm,
        title=ctx["document_title"],
        author="Zarrah Foundation",
        subject=f"{ctx['resolution_id']} v{ctx['version_number']}",
        creator="Zarrah Foundation Governance",
    )

    title_style = ParagraphStyle(
        "ResolutionDocTitle",
        parent=styles["heading"],
        fontSize=16,
        leading=20,
        spaceAfter=4,
        shaping=1,
    )
    subtitle_style = ParagraphStyle(
        "ResolutionDocSubtitle",
        parent=styles["body"],
        fontSize=11,
        leading=14,
        textColor=colors.HexColor("#333333"),
        spaceAfter=8,
        shaping=1,
    )
    section_style = ParagraphStyle(
        "ResolutionSection",
        parent=styles["heading"],
        fontSize=12,
        leading=16,
        spaceBefore=8,
        spaceAfter=6,
        shaping=1,
    )
    body_style = styles["body"]
    label_style = styles["label"]
    small_style = styles["small"]

    story: list[Any] = []

    if ctx["logo_path"]:
        story.append(
            Image(
                str(ctx["logo_path"]),
                width=32 * mm,
                height=32 * mm,
                kind="proportional",
            )
        )
        story.append(Spacer(1, 3 * mm))

    story.append(
        Paragraph(escape_pdf_text(ctx["foundation_name"], bold=True), title_style)
    )
    if (
        ctx["foundation_name_alt"]
        and ctx["foundation_name_alt"] != ctx["foundation_name"]
    ):
        story.append(Paragraph(escape_pdf_text(ctx["foundation_name_alt"]), subtitle_style))
    if ctx.get("registration_number"):
        story.append(
            Paragraph(
                escape_pdf_text(
                    _("Registration No.: %(num)s")
                    % {"num": ctx["registration_number"]}
                ),
                small_style,
            )
        )
    story.append(Paragraph(escape_pdf_text(_("Resolution Document")), subtitle_style))
    story.append(Spacer(1, 4 * mm))

    meta_rows = [
        [_("Resolution ID"), ctx["resolution_id"]],
        [_("Title"), ctx["resolution_title"]],
        [_("Meeting"), f"{ctx['meeting_reference']} — {ctx['meeting_title']}"],
        [_("Meeting type"), ctx["meeting_type"]],
        [_("Date & time"), ctx["meeting_starts_at_display"]],
        [_("Mode"), ctx["meeting_mode"]],
    ]
    if ctx["meeting_venue"]:
        meta_rows.append([_("Venue"), ctx["meeting_venue"]])
    if ctx["agenda_title"]:
        agenda_label = _("Agenda item")
        if ctx["agenda_number"] is not None:
            agenda_label = f"{agenda_label} {ctx['agenda_number']:02d}"
        meta_rows.append([agenda_label, ctx["agenda_title"]])
    meta_rows.extend(
        [
            [_("Version"), str(ctx["version_number"])],
            [_("Status"), ctx["version_status_line"]],
        ]
    )

    meta_table_data = [
        [
            Paragraph(f"<b>{escape_pdf_text(label)}</b>", label_style),
            Paragraph(escape_pdf_text(value), body_style),
        ]
        for label, value in meta_rows
    ]
    meta_table = Table(meta_table_data, colWidths=[48 * mm, 124 * mm])
    meta_table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#F3F7F4")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 8),
                ("RIGHTPADDING", (0, 0), (-1, -1), 8),
                ("TOPPADDING", (0, 0), (-1, -1), 6),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
                ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor("#D0D7D2")),
                ("INNERGRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#E2E8E4")),
            ]
        )
    )
    story.append(meta_table)
    story.append(Spacer(1, 6 * mm))

    story.append(Paragraph(escape_pdf_text(_("Resolution text"), bold=True), section_style))
    if ctx["content_paragraphs_bn"]:
        story.append(Paragraph(escape_pdf_text(_("Bangla"), bold=True), label_style))
        for paragraph in ctx["content_paragraphs_bn"]:
            story.append(Paragraph(escape_pdf_text(paragraph), body_style))
            story.append(Spacer(1, 2 * mm))
    if ctx["content_paragraphs_en"]:
        story.append(Paragraph(escape_pdf_text(_("English"), bold=True), label_style))
        for paragraph in ctx["content_paragraphs_en"]:
            story.append(Paragraph(escape_pdf_text(paragraph), body_style))
            story.append(Spacer(1, 2 * mm))
    if not ctx["content_paragraphs_bn"] and not ctx["content_paragraphs_en"]:
        story.append(Paragraph(escape_pdf_text("—"), body_style))

    if ctx["decision_paragraphs_bn"] or ctx["decision_paragraphs_en"]:
        story.append(Spacer(1, 4 * mm))
        story.append(
            Paragraph(escape_pdf_text(_("Decision"), bold=True), section_style)
        )
        if ctx["decision_paragraphs_bn"]:
            for paragraph in ctx["decision_paragraphs_bn"]:
                story.append(Paragraph(escape_pdf_text(paragraph), body_style))
                story.append(Spacer(1, 2 * mm))
        if ctx["decision_paragraphs_en"]:
            for paragraph in ctx["decision_paragraphs_en"]:
                story.append(Paragraph(escape_pdf_text(paragraph), body_style))
                story.append(Spacer(1, 2 * mm))

    story.append(Spacer(1, 4 * mm))
    story.append(
        Paragraph(
            escape_pdf_text(
                _(
                    "This document records the decisions of the meeting for "
                    "authorized Foundation administration."
                )
            ),
            small_style,
        )
    )

    signatures = ctx["signatures"]
    if signatures:
        story.append(Spacer(1, 6 * mm))
        story.append(
            Paragraph(escape_pdf_text(_("Signatures"), bold=True), section_style)
        )
        sig_cells: list[list[Any]] = []
        row: list[Any] = []
        for sig in signatures:
            block: list[Any] = []
            if sig["is_signed"]:
                block.append(
                    Paragraph(
                        escape_pdf_text(_("Approved & Signed"), bold=True),
                        small_style,
                    )
                )
            if sig["image_path"]:
                block.append(
                    Image(
                        str(sig["image_path"]),
                        width=45 * mm,
                        height=18 * mm,
                        kind="proportional",
                    )
                )
                block.append(Spacer(1, 2 * mm))
            block.append(
                Paragraph(
                    escape_pdf_text(sig["name"], bold=True),
                    ParagraphStyle(
                        "SigName",
                        parent=body_style,
                        fontSize=10,
                        leading=13,
                        shaping=1,
                    ),
                )
            )
            if sig["designation"]:
                block.append(Paragraph(escape_pdf_text(sig["designation"]), small_style))
            block.append(Paragraph(escape_pdf_text(sig["member_id"]), small_style))
            if sig["is_signed"]:
                block.append(
                    Paragraph(
                        escape_pdf_text(
                            _("Signed: %(date)s") % {"date": sig["signed_at_display"]}
                        ),
                        small_style,
                    )
                )
            else:
                block.append(
                    Paragraph(escape_pdf_text(_("Awaiting signature")), small_style)
                )
            row.append(block)
            if len(row) == 2:
                sig_cells.append(row)
                row = []
        if row:
            row.append(Spacer(1, 1))
            sig_cells.append(row)

        sig_table = Table(sig_cells, colWidths=[86 * mm, 86 * mm])
        sig_table.setStyle(
            TableStyle(
                [
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("LEFTPADDING", (0, 0), (-1, -1), 4),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 4),
                    ("TOPPADDING", (0, 0), (-1, -1), 6),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 10),
                    ("BOX", (0, 0), (-1, -1), 0.25, colors.HexColor("#D0D7D2")),
                ]
            )
        )
        story.append(sig_table)

    if ctx.get("website_url"):
        story.append(Spacer(1, 6 * mm))
        story.append(Paragraph(escape_pdf_text(ctx["website_url"]), small_style))

    doc.build(story, onFirstPage=_draw_page_number, onLaterPages=_draw_page_number)
    if actor is not None:
        _log_document_generation(
            resolution,
            actor=actor,
            fmt="pdf",
            version_number=ctx["version_number"],
        )
    return buffer.getvalue()


def _docx_add_page_number(paragraph) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_sep)
    run._r.append(fld_end)


def generate_resolution_docx(resolution, *, actor=None) -> bytes:
    """Render an official resolution DOCX with logo, signatures, and page numbers."""
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Mm, Pt

    ctx = build_resolution_document_context(resolution)
    document = Document()
    props = document.core_properties
    props.author = "Zarrah Foundation"
    props.title = str(ctx["document_title"])
    props.subject = f"{ctx['resolution_id']} v{ctx['version_number']}"
    props.category = "Resolution"
    props.language = ctx["language"]

    section = document.sections[0]
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer.add_run(_("Page "))
    footer_run.font.size = Pt(8)
    _docx_add_page_number(footer)
    footer.add_run(f"  ·  {ctx['generated_at_display']}").font.size = Pt(8)

    if ctx["logo_path"]:
        document.add_picture(str(ctx["logo_path"]), width=Mm(32))
        document.add_paragraph()

    title = document.add_heading(ctx["foundation_name"], level=1)
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    if (
        ctx["foundation_name_alt"]
        and ctx["foundation_name_alt"] != ctx["foundation_name"]
    ):
        alt = document.add_paragraph(ctx["foundation_name_alt"])
        alt.runs[0].font.size = Pt(11)
    if ctx.get("registration_number"):
        reg = document.add_paragraph(
            _("Registration No.: %(num)s") % {"num": ctx["registration_number"]}
        )
        reg.runs[0].font.size = Pt(9)
    document.add_paragraph(_("Resolution Document")).runs[0].italic = True
    document.add_paragraph()

    meta = document.add_table(rows=0, cols=2)
    meta.style = "Table Grid"
    meta_rows = [
        (_("Resolution ID"), ctx["resolution_id"]),
        (_("Title"), ctx["resolution_title"]),
        (_("Meeting"), f"{ctx['meeting_reference']} — {ctx['meeting_title']}"),
        (_("Meeting type"), ctx["meeting_type"]),
        (_("Date & time"), ctx["meeting_starts_at_display"]),
        (_("Mode"), ctx["meeting_mode"]),
    ]
    if ctx["meeting_venue"]:
        meta_rows.append((_("Venue"), ctx["meeting_venue"]))
    if ctx["agenda_title"]:
        agenda_label = _("Agenda item")
        if ctx["agenda_number"] is not None:
            agenda_label = f"{agenda_label} {ctx['agenda_number']:02d}"
        meta_rows.append((agenda_label, ctx["agenda_title"]))
    meta_rows.extend(
        [
            (_("Version"), str(ctx["version_number"])),
            (_("Status"), ctx["version_status_line"]),
        ]
    )
    for label, value in meta_rows:
        row = meta.add_row()
        row.cells[0].text = str(label)
        row.cells[1].text = str(value)
        for run in row.cells[0].paragraphs[0].runs:
            run.bold = True
            run.font.name = "Noto Sans Bengali"
        for run in row.cells[1].paragraphs[0].runs:
            run.font.name = "Noto Sans Bengali"

    document.add_paragraph()
    document.add_heading(_("Resolution text"), level=2)
    if ctx["content_paragraphs_bn"]:
        document.add_paragraph(_("Bangla")).runs[0].bold = True
        for paragraph in ctx["content_paragraphs_bn"]:
            p = document.add_paragraph(paragraph)
            for run in p.runs:
                run.font.name = "Noto Sans Bengali"
    if ctx["content_paragraphs_en"]:
        document.add_paragraph(_("English")).runs[0].bold = True
        for paragraph in ctx["content_paragraphs_en"]:
            p = document.add_paragraph(paragraph)
            for run in p.runs:
                run.font.name = "Noto Sans Bengali"
    if not ctx["content_paragraphs_bn"] and not ctx["content_paragraphs_en"]:
        document.add_paragraph("—")

    if ctx["decision_paragraphs_bn"] or ctx["decision_paragraphs_en"]:
        document.add_heading(_("Decision"), level=2)
        for paragraph in ctx["decision_paragraphs_bn"]:
            p = document.add_paragraph(paragraph)
            for run in p.runs:
                run.font.name = "Noto Sans Bengali"
        for paragraph in ctx["decision_paragraphs_en"]:
            p = document.add_paragraph(paragraph)
            for run in p.runs:
                run.font.name = "Noto Sans Bengali"

    document.add_paragraph(
        _(
            "This document records the decisions of the meeting for "
            "authorized Foundation administration."
        )
    )

    if ctx["signatures"]:
        document.add_heading(_("Signatures"), level=2)
        sig_table = document.add_table(rows=0, cols=2)
        sig_table.style = "Table Grid"
        signatures = ctx["signatures"]
        for index in range(0, len(signatures), 2):
            row = sig_table.add_row()
            for col in range(2):
                if index + col >= len(signatures):
                    continue
                sig = signatures[index + col]
                cell = row.cells[col]
                cell.text = ""
                p = cell.paragraphs[0]
                if sig["is_signed"]:
                    approved = p.add_run(_("Approved & Signed") + "\n")
                    approved.bold = True
                if sig["image_path"]:
                    cell.add_paragraph()
                    cell.paragraphs[-1].add_run().add_picture(
                        str(sig["image_path"]),
                        width=Mm(45),
                    )
                    p = cell.add_paragraph()
                name_run = p.add_run(sig["name"])
                name_run.bold = True
                name_run.font.name = "Noto Sans Bengali"
                if sig["designation"]:
                    p.add_run(f"\n{sig['designation']}")
                p.add_run(f"\n{sig['member_id']}")
                if sig["is_signed"]:
                    p.add_run(
                        f"\n{_('Signed: %(date)s') % {'date': sig['signed_at_display']}}"
                    )
                else:
                    p.add_run(f"\n{_('Awaiting signature')}")

    if ctx.get("website_url"):
        document.add_paragraph(ctx["website_url"])

    buffer = io.BytesIO()
    document.save(buffer)
    if actor is not None:
        _log_document_generation(
            resolution,
            actor=actor,
            fmt="docx",
            version_number=ctx["version_number"],
        )
    return buffer.getvalue()


def resolution_document_filename(resolution, *, ext: str = "pdf", version=None) -> str:
    """Suggested download filename: ``ZF-2026-07-001-v1.pdf``."""
    ext = ext.lstrip(".").lower() or "pdf"
    rid = re.sub(r"[^A-Za-z0-9\-]+", "-", resolution.resolution_id).strip("-")
    ver = version
    if ver is None:
        ver = getattr(getattr(resolution, "current_version", None), "version_number", 1)
    return f"{rid}-v{int(ver)}.{ext}"


def assert_no_protected_media_urls(payload: bytes) -> None:
    """Raise ValueError if generated bytes embed protected-media URL paths."""
    text = payload.decode("latin-1", errors="ignore").lower()
    needles = (
        "protected-media",
        "/media/members/signatures/",
        "/media/resolutions/signatures/",
    )
    if any(n in text for n in needles):
        raise ValueError("Document must not embed private media URLs.")


def resolution_file_response(resolution, *, fmt: str, actor=None):
    """Build an HttpResponse for on-demand PDF/DOCX download (no disk cache)."""
    from django.http import HttpResponse

    fmt = (fmt or "pdf").lower().lstrip(".")
    if fmt == "pdf":
        payload = generate_resolution_pdf(resolution, actor=actor)
        content_type = "application/pdf"
    elif fmt == "docx":
        payload = generate_resolution_docx(resolution, actor=actor)
        content_type = (
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        )
    else:
        raise ValueError(f"Unsupported document format: {fmt}")
    assert_no_protected_media_urls(payload)
    filename = resolution_document_filename(resolution, ext=fmt)
    response = HttpResponse(payload, content_type=content_type)
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    response["X-Content-Type-Options"] = "nosniff"
    return response
